from langchain_core.documents import Document
from pathlib import Path
import requests
from bs4 import BeautifulSoup
import os
import logging

logger = logging.getLogger(__name__)


def load_documents(source: str) -> list[Document]:
    """Load documents from a URL or local file with comprehensive error handling."""
    if not source or not str(source).strip():
        raise ValueError("Source path or URL cannot be empty")

    source_str = str(source).strip()

    if source_str.startswith("http://") or source_str.startswith("https://"):
        return _load_url(source_str)

    path = Path(source_str)

    if path.exists() and path.is_file():
        if path.stat().st_size == 0:
            raise ValueError(f"File '{path.name}' is empty (0 bytes).")

        ext = path.suffix.lower()

        if ext == ".pdf":
            logger.info(f"Loading PDF: {source_str}")
            docs = _load_pdf(str(path))

        elif ext == ".txt":
            logger.info(f"Loading TXT: {source_str}")
            docs = _load_txt(str(path))

        elif ext == ".docx":
            logger.info(f"Loading DOCX: {source_str}")
            docs = _load_docx(str(path))

        else:
            raise ValueError(f"Unsupported file type: '{ext}'. Supported types are .pdf, .txt, .docx")

    else:
        logger.info("Loading raw text input")
        docs = [Document(page_content=source_str, metadata={"source": "raw_text"})]

    # Filter empty docs
    docs = [d for d in docs if d.page_content.strip()]
    if not docs:
        raise ValueError("Document contains no readable text content.")

    logger.info(f"Loaded {len(docs)} non-empty document(s)")
    return docs


def _load_pdf(path: str) -> list[Document]:
    """
    Extract text from PDF using PyMuPDF (fitz) or pypdf fallback
    with corruption, password protection, and scanned-PDF handling.
    """
    pages_text = []

    # Attempt 1: PyMuPDF (fitz)
    use_fitz = True
    try:
        import fitz
        doc = fitz.open(path)
        try:
            if doc.is_encrypted:
                raise ValueError("This PDF is password-protected and cannot be read.")
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text").strip()
                if text:
                    pages_text.append(text)
        finally:
            doc.close()
    except (ImportError, Exception) as fitz_err:
        logger.info(f"PyMuPDF unavailable or failed ({fitz_err}), attempting pypdf fallback...")
        use_fitz = False

    # Attempt 2: pypdf fallback
    if not use_fitz or not pages_text:
        try:
            import pypdf
            reader = pypdf.PdfReader(path)
            if reader.is_encrypted:
                raise ValueError("This PDF is password-protected and cannot be read.")
            for page in reader.pages:
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(text.strip())
        except ValueError:
            raise
        except Exception as pypdf_err:
            raise ValueError(f"Failed to parse PDF file with pypdf: {pypdf_err}")

    if not pages_text:
        raise ValueError(
            "This PDF contains no extractable text. It appears to be a scanned or image-only PDF. "
            "Please run OCR on it or copy the text into a .txt file before uploading."
        )

    full_text = "\n\n".join(pages_text)
    logger.info(f"Extracted {len(full_text):,} characters from {len(pages_text)} pages in PDF")

    return [Document(
        page_content=full_text,
        metadata={"source": path, "type": "pdf", "pages": len(pages_text)}
    )]


def _load_txt(path: str) -> list[Document]:
    """Load plain text file supporting multiple character encodings."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    content = None

    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                content = f.read()
            break
        except UnicodeDecodeError:
            continue

    if content is None:
        # Fallback with replacement characters
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()

    if not content.strip():
        raise ValueError("Text file appears to be empty or contains only whitespace.")

    return [Document(
        page_content=content,
        metadata={"source": path, "type": "txt"}
    )]


def _load_docx(path: str) -> list[Document]:
    """Extract text from DOCX with error handling."""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(path)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document (file may be damaged or invalid): {e}")

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)
    if not full_text.strip():
        raise ValueError("DOCX file contains no readable text content.")

    return [Document(
        page_content=full_text,
        metadata={"source": path, "type": "docx"}
    )]


def _load_url(url: str) -> list[Document]:
    """Fetch and extract readable content from a web URL."""
    logger.info(f"Fetching URL: {url}")
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
    }

    try:
        response = requests.get(url, headers=headers, timeout=20)
        if response.status_code == 403:
            raise ValueError(f"Access forbidden (HTTP 403). The website '{url}' blocks automated requests.")
        if response.status_code == 404:
            raise ValueError(f"Page not found (HTTP 404) at '{url}'.")
        response.raise_for_status()
    except requests.Timeout:
        raise ValueError(f"Connection timed out while trying to reach '{url}'.")
    except requests.ConnectionError:
        raise ValueError(f"Could not connect to '{url}'. Please verify the domain and your network connection.")
    except requests.RequestException as e:
        raise ValueError(f"Failed to fetch URL '{url}': {e}")

    content_type = response.headers.get("content-type", "").lower()
    if "text" not in content_type and "html" not in content_type:
        raise ValueError(f"URL did not return text or HTML content (Content-Type: {content_type}).")

    try:
        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header",
                         "aside", "iframe", "noscript", "svg"]):
            tag.decompose()

        text = soup.get_text(separator="\n", strip=True)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        clean_text = "\n\n".join(lines)
    except Exception as e:
        raise ValueError(f"Failed to parse HTML from '{url}': {e}")

    if len(clean_text) < 50:
        raise ValueError(
            f"Could not extract meaningful text from '{url}'. The page might require JavaScript rendering or login."
        )

    title = soup.find("title")
    title_text = title.get_text(strip=True) if title else url
    logger.info(f"Extracted {len(clean_text):,} chars from: {title_text}")

    return [Document(
        page_content=clean_text,
        metadata={"source": url, "title": title_text, "type": "url"}
    )]